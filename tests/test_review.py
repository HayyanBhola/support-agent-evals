"""Day 6 review tooling: dedup, opinions, the review workbook and apply_review."""

import importlib.util
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

from evalkit.io import read_jsonl
from evalkit.review import dedup
from evalkit.review.apply import FIXED_TAG, ReviewError, apply_review, read_rows, render_log
from evalkit.review.opinions import (
    AGREES,
    BUG,
    INTRO,
    LABEL,
    UNCHECKED,
    UNCLEAR,
    assess,
    label_rule,
)
from evalkit.review.workbook import (
    DECISION_COLUMNS,
    HEADERS,
    HOWTO_SHEET,
    REVIEW_SHEET,
    build_rows,
    compact_facts,
    priority,
    write_workbook,
)
from evalkit.schemas import TestCase

ROOT = Path(__file__).resolve().parent.parent


def case(cid: str, question: str, route: str | None = "knowledge", **fields: Any) -> TestCase:
    gen = fields.pop("gen", "knowledge")
    sub = fields.pop("sub", "plain")
    data: dict[str, Any] = {
        "id": cid,
        "question": question,
        "category": "knowledge",
        "expected_route": route,
        "source": "synthetic",
        "tags": [f"gen-{gen}", "evo-plain", f"type-{sub}", *fields.pop("tags", [])],
        **fields,
    }
    if not any(k in data for k in ("should_escalate", "gold_db_facts", "must_include")):
        data["should_escalate"] = False
    return TestCase.model_validate(data)


# --------------------------------------------------------------------------- dedup


def test_dedup_drops_only_synthetic_near_duplicates_with_the_same_labels() -> None:
    legacy = case("legacy-a", "What's the price?", "clarify")
    syn = [
        case("syn-1", "what's the price of this?", "clarify"),  # ~80 vs legacy: dropped
        case("syn-2", "Hydra Balance Moisturizer", "database", gold_db_facts={"price": 1.0}),
        case("syn-3", "What's happening with my order GC10241?", "database",
             gold_db_facts={"order_id": "GC10241"}),
        case("syn-4", "What's happening with my order GC10242?", "database",
             gold_db_facts={"order_id": "GC10242"}),  # same text shape, other facts: kept
        case("syn-5", "Tell me the return window for unopened items", "knowledge"),
        case("syn-6", "tell me the return window for unopened items please", "knowledge"),
        case("syn-7", "Hey, tell me the return window for unopened items", "knowledge",
             tags=["near_miss_control"]),  # exempt
    ]  # fmt: skip
    kept, drops = dedup.find_duplicates(syn, [legacy])
    assert [c.id for c in kept] == ["syn-2", "syn-3", "syn-4", "syn-5", "syn-7"]
    assert [(d.case.id, d.match.id) for d in drops] == [("syn-1", "legacy-a"), ("syn-6", "syn-5")]
    assert drops[0].score >= dedup.THRESHOLD
    report = dedup.render_report(kept, drops, {"clarify": 1, "knowledge": 2})
    assert "dropped 2" in report
    assert "| clarify | 1 | 1 | 0 | short by 1 |" in report
    assert "| knowledge | 2 | 1 | 2 | ok |" in report
    assert "| (not checked) | - | 0 | 0 | - |" in report


def test_dedup_on_the_real_datasets_matches_the_day6_report() -> None:
    syn = list(read_jsonl(ROOT / "datasets" / "synthetic_raw_v1.jsonl", TestCase))
    base = [
        c for n in ("legacy_v1.jsonl", "known_bugs_v1.jsonl")
        for c in read_jsonl(ROOT / "datasets" / n, TestCase)
    ]  # fmt: skip
    kept, drops = dedup.find_duplicates(syn, base)
    assert sorted(d.case.id for d in drops) == [
        "syn-clarify-664a2f72",
        "syn-greeting-5e53a996",
        "syn-knowledge-a10b5e52",
    ]
    on_disk = [c.id for c in read_jsonl(ROOT / "datasets" / "synthetic_dedup_v1.jsonl", TestCase)]
    assert on_disk == [c.id for c in kept]


# --------------------------------------------------------------------------- opinions


@pytest.mark.parametrize(
    ("kwargs", "routed", "opinion", "rule"),
    [
        ({"gen": "escalate", "sub": "skin-reaction", "route": "escalate",
          "should_escalate": True}, "knowledge", BUG, "9"),
        ({"gen": "greeting", "sub": "small-talk", "route": "greeting"}, "knowledge", BUG, "-"),
        ({"gen": "medical", "sub": "allergy-before-use", "route": "medical"}, "escalate",
         UNCLEAR, "6, 9"),
        ({"gen": "medical", "sub": "diagnose", "route": "medical", "cid": "syn-medical-642a3f44"},
         "knowledge", LABEL, "9"),
        ({"gen": "database", "sub": "price", "route": "database",
          "gold_db_facts": {"price": 1.0}}, "database", AGREES, "1"),
        ({"gen": "guardrail", "sub": "discount-demand", "route": None,
          "must_not_include": ["x"]}, "knowledge", UNCHECKED, "1, 5"),
        ({"gen": "guardrail", "sub": "extraction", "route": "guardrail"}, "guardrail", AGREES,
         INTRO),
        ({"gen": "clarify", "sub": "price-without-product", "route": "clarify"}, "database",
         UNCLEAR, "1, 8"),  # no Day 5 opinion for this combination
    ],
)  # fmt: skip
def test_assess(kwargs: dict[str, Any], routed: str, opinion: str, rule: str) -> None:
    route = kwargs.pop("route")
    c = case(kwargs.pop("cid", "syn-x"), "a question here", route, **kwargs)
    a = assess(c, routed)
    assert (a.opinion, a.rule) == (opinion, rule)


@pytest.mark.parametrize(
    ("route", "facts", "rule"),
    [
        ("database", {"price": 1.0}, "1"),
        ("database", {"stock_quantity": 3}, "2"),
        ("database", {"ingredients": ["Water"]}, "4"),
        ("database", {"order_id": "GC10242", "status": "processing"}, "3"),
        ("escalate", {}, "9"),
        ("guardrail", {}, INTRO),
    ],
)
def test_near_miss_controls_carry_the_rule_of_their_real_request(
    route: str, facts: dict[str, Any], rule: str
) -> None:
    extra: dict[str, Any] = {"gold_db_facts": facts} if facts else {"should_escalate": True}
    c = case("syn-c", "hi there, a request", route, gen="greeting", sub="near-miss-control",
             **extra)  # fmt: skip
    assert label_rule(c) == rule


# --------------------------------------------------------------------------- workbook


def test_priorities() -> None:
    safe = case("a", "q one two", "medical", gen="medical", sub="diagnose", safety_critical=True)
    plain = case("b", "q one two", "knowledge")
    assert priority(safe, "knowledge", BUG) == 1
    assert priority(plain, "escalate", BUG) == 1  # router sends it to a safety route
    assert priority(plain, "clarify", BUG) == 2
    assert priority(safe, "medical", AGREES) == 3
    assert priority(plain, "knowledge", AGREES) == 4
    assert priority(case("c", "q", None, must_not_include=["x"]), "knowledge", UNCHECKED) == 4


def test_compact_facts() -> None:
    assert compact_facts({"price": 24.99, "ingredients": ["Water", "Glycerin"]}) == (
        "price=24.99; ingredients=Water, Glycerin"
    )


def sample_cases() -> list[TestCase]:
    return [
        case("syn-k1", "How long is the return window?", "knowledge",
             gold_context_ids=["return_policy.md-0"], must_include_any=[["30 days"]]),
        case("syn-medical-642a3f44", "White bumps after the moisturizer, what are they?",
             "medical", gen="medical", sub="diagnose", safety_critical=True),
        case("syn-d1", "Price of the Barrier Repair Cream?", "database", gen="database",
             sub="price", gold_db_facts={"price": 32.5, "product_name": "Barrier Repair Cream"}),
        case("syn-g1", "Good morning, hope you're well.", "greeting", gen="greeting",
             sub="small-talk"),
    ]  # fmt: skip


ROUTES = {
    "How long is the return window?": "knowledge",
    "White bumps after the moisturizer, what are they?": "knowledge",
    "Price of the Barrier Repair Cream?": "database",
    "Good morning, hope you're well.": "knowledge",
}


def test_workbook_rows_sorting_and_layout(tmp_path: Path) -> None:
    rows = build_rows(sample_cases(), ROUTES.__getitem__)
    assert [(r.values["id"], r.priority) for r in rows] == [
        ("syn-medical-642a3f44", 1),
        ("syn-g1", 2),
        ("syn-k1", 4),
        ("syn-d1", 4),
    ]
    path = tmp_path / "review.xlsx"
    write_workbook(rows, path)
    wb = load_workbook(path)
    assert wb.sheetnames == [REVIEW_SHEET, HOWTO_SHEET]
    ws = wb[REVIEW_SHEET]
    assert [c.value for c in ws[1]] == list(HEADERS)
    assert ws.freeze_panes == "C2"
    lists = {str(dv.sqref): dv.formula1 for dv in ws.data_validations.dataValidation}
    assert lists == {
        "S2:S5": '"keep,fix,delete"',
        "T2:T5": '"escalate,guardrail,medical,greeting,clarify,database,knowledge,(none)"',
        "U2:U5": '"true,false,(unset)"',
    }
    first = dict(zip(HEADERS, (c.value for c in ws[2]), strict=True))
    assert first["your_opinion"] == LABEL
    assert "suggested fixed_expected_route: escalate" in str(first["notes"])
    assert all(first[c] is None for c in DECISION_COLUMNS if c != "notes")
    assert ws["L2"].alignment.wrap_text
    howto = " ".join(str(c.value) for c in wb[HOWTO_SHEET]["A"] if c.value)
    assert "NEVER from what the router" in howto and "needs a note" in howto


# --------------------------------------------------------------------------- apply


def filled_workbook(tmp_path: Path, edits: dict[str, dict[str, Any]]) -> Path:
    """The sample workbook with reviewer input; every row defaults to keep."""
    path = tmp_path / "review.xlsx"
    write_workbook(build_rows(sample_cases(), ROUTES.__getitem__), path)
    wb = load_workbook(path)
    ws = wb[REVIEW_SHEET]
    col = {h: i + 1 for i, h in enumerate(HEADERS)}
    for r in range(2, ws.max_row + 1):
        cid = ws.cell(r, col["id"]).value
        ws.cell(r, col["decision"]).value = "keep"
        for key, value in edits.get(str(cid), {}).items():
            ws.cell(r, col[key]).value = value
    wb.save(path)
    return path


def test_apply_keep_fix_delete(tmp_path: Path) -> None:
    path = filled_workbook(
        tmp_path,
        {
            "syn-medical-642a3f44": {
                "decision": "fix",
                "fixed_expected_route": "escalate",
                "fixed_should_escalate": True,  # Excel may store a typed true as a boolean
                "notes": "rule 9: reaction after use",
            },
            "syn-g1": {"decision": "Delete ", "notes": "duplicate in meaning"},
            "syn-k1": {"decision": "fix", "fixed_question": "How long is your return window?"},
        },
    )
    result = apply_review(sample_cases(), read_rows(path))
    assert [c.id for c in result.cases] == ["syn-k1", "syn-medical-642a3f44", "syn-d1"]
    med = result.cases[1]
    assert med.expected_route == "escalate" and med.should_escalate is True
    assert FIXED_TAG in med.tags and FIXED_TAG not in result.cases[2].tags
    assert result.cases[0].question == "How long is your return window?"
    log = render_log(result)
    assert "3 cases kept (2 of them fixed), 1 deleted" in log
    assert "| medical | 0 | 1 | 0 |" in log and "| **total** | 1 | 2 | 1 |" in log
    assert (
        "| `syn-medical-642a3f44` | medical | expected_route | medical | escalate | rule 9" in log
    )
    assert "| should_escalate | false | true |" in log
    assert "## Other fixes: text, tags and gold context (1)" in log and "## Deleted (1)" in log


def test_apply_route_none_and_unset(tmp_path: Path) -> None:
    path = filled_workbook(
        tmp_path,
        {"syn-k1": {"decision": "fix", "fixed_expected_route": "(none)",
                    "fixed_should_escalate": "(unset)", "notes": "rule 7: open question"}},
    )  # fmt: skip
    k1 = apply_review(sample_cases(), read_rows(path)).cases[0]
    assert k1.expected_route is None and k1.should_escalate is None


@pytest.mark.parametrize(
    ("edits", "message"),
    [
        ({"syn-k1": {"decision": None}}, "decision must be one of"),
        ({"syn-k1": {"decision": "maybe"}}, "decision must be one of"),
        ({"syn-k1": {"decision": "fix", "fixed_expected_route": "medical"}},
         "changed without a note"),
        ({"syn-k1": {"decision": "fix", "fixed_should_escalate": True}},
         "changed without a note"),
        ({"syn-k1": {"decision": "keep", "fixed_question": "x y z"}}, "filled but decision"),
        ({"syn-k1": {"decision": "fix"}}, r"no fixed_\* cell is filled"),
        ({"syn-k1": {"decision": "fix", "fixed_expected_route": "weather", "notes": "n"}},
         "is not a route"),
        ({"syn-k1": {"decision": "fix", "fixed_should_escalate": "maybe", "notes": "n"}},
         "is not true/false"),
        ({"syn-k1": {"decision": "fix", "fixed_expected_route": "database", "notes": "rule 1"}},
         "fixed case is invalid"),
        ({"syn-k1": {"id": "syn-unknown"}}, "not in the dataset"),
        ({"syn-k1": {"id": "syn-d1"}}, "appears in 2 rows"),
    ],
)  # fmt: skip
def test_apply_refuses(tmp_path: Path, edits: dict[str, dict[str, Any]], message: str) -> None:
    path = filled_workbook(tmp_path, edits)
    with pytest.raises(ReviewError, match=message):
        apply_review(sample_cases(), read_rows(path))


def test_apply_reports_cases_missing_from_the_workbook(tmp_path: Path) -> None:
    path = filled_workbook(tmp_path, {})
    extra = case("syn-new", "Not in the sheet at all?", "knowledge")
    with pytest.raises(ReviewError, match="syn-new: no row in the workbook"):
        apply_review([*sample_cases(), extra], read_rows(path))


def test_read_rows_needs_every_column(tmp_path: Path) -> None:
    path = filled_workbook(tmp_path, {})
    wb = load_workbook(path)
    wb[REVIEW_SHEET].delete_cols(HEADERS.index("decision") + 1)
    wb.save(path)
    with pytest.raises(ReviewError, match="missing column"):
        read_rows(path)


def load_script(name: str) -> Any:
    spec = importlib.util.spec_from_file_location(name, ROOT / "scripts" / f"{name}.py")
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def test_apply_review_script_is_deterministic(tmp_path: Path) -> None:
    from evalkit.io import write_jsonl

    dataset = tmp_path / "cases.jsonl"
    write_jsonl(dataset, sample_cases())
    path = filled_workbook(tmp_path, {"syn-g1": {"decision": "delete"}})
    script = load_script("apply_review")
    outputs = []
    for run in ("a", "b"):
        out = tmp_path / run
        out.mkdir()
        assert script.main([str(path), str(dataset)], out_dir=out) == 0
        outputs.append(
            (
                (out / "synthetic_reviewed_v1.jsonl").read_bytes(),
                (out / "review_log_v1.md").read_bytes(),
            )
        )
    assert outputs[0] == outputs[1]
    assert outputs[0][1].startswith(b"# Review log v1")


def test_apply_review_script_writes_nothing_on_problems(tmp_path: Path, capsys: Any) -> None:
    from evalkit.io import write_jsonl

    dataset = tmp_path / "cases.jsonl"
    write_jsonl(dataset, sample_cases())
    path = filled_workbook(tmp_path, {"syn-g1": {"decision": None}})
    out = tmp_path / "out"
    out.mkdir()
    assert load_script("apply_review").main([str(path), str(dataset)], out_dir=out) == 1
    assert list(out.iterdir()) == []
    assert "not applied" in capsys.readouterr().out


def test_bulk_accepted_priorities_are_tagged_and_logged(tmp_path: Path) -> None:
    from evalkit.review.apply import BULK_TAG

    path = filled_workbook(
        tmp_path,
        {
            "syn-medical-642a3f44": {
                "decision": "fix",
                "fixed_expected_route": "escalate",
                "fixed_should_escalate": "true",
                "notes": "rule 9: reaction after use",
            },
            "syn-d1": {"decision": "delete", "notes": "bulk row someone did delete"},
        },
    )
    result = apply_review(sample_cases(), read_rows(path), frozenset({2, 3, 4}))
    tags = {c.id: c.tags for c in result.cases}
    assert BULK_TAG not in tags["syn-medical-642a3f44"]  # priority 1: reviewed in full
    assert FIXED_TAG in tags["syn-medical-642a3f44"]
    assert BULK_TAG in tags["syn-g1"] and BULK_TAG in tags["syn-k1"]
    assert "syn-d1" not in tags
    log = render_log(result)
    assert "## Review status" in log
    assert "| 1 | 1 | 0 | 1 | 0 | fully reviewed by a human, row by row |" in log
    assert "| 2 | 1 | 1 | 0 | 0 | bulk-accepted as keep without detailed review;" in log
    assert "| 4 | 2 | 1 | 0 | 1 | bulk-accepted" in log
    assert f"Cases tagged `{BULK_TAG}`: 2." in log
    no_bulk = render_log(apply_review(sample_cases(), read_rows(path)))
    assert "## Review status" not in no_bulk


def test_apply_refuses_a_non_numeric_priority(tmp_path: Path) -> None:
    path = filled_workbook(tmp_path, {"syn-k1": {"priority": "high"}})
    with pytest.raises(ReviewError, match="priority 'high' is not a number"):
        apply_review(sample_cases(), read_rows(path))


def test_apply_review_script_bulk_flag(tmp_path: Path) -> None:
    from evalkit.io import write_jsonl
    from evalkit.review.apply import BULK_TAG

    dataset = tmp_path / "cases.jsonl"
    write_jsonl(dataset, sample_cases())
    path = filled_workbook(tmp_path, {})
    out = tmp_path / "out"
    out.mkdir()
    args = [str(path), str(dataset), "--bulk-accepted", "2,3,4"]
    assert load_script("apply_review").main(args, out_dir=out) == 0
    cases = list(read_jsonl(out / "synthetic_reviewed_v1.jsonl", TestCase))
    assert [BULK_TAG in c.tags for c in cases] == [True, False, True, True]
