"""Apply a filled-in review workbook to the synthetic cases.

Rules (the run is refused, nothing written, if any is broken):

* every case in the dataset has exactly one row, and every row's id is in the dataset;
* every row has a decision: keep, fix or delete;
* only ``fix`` rows may fill fixed_* cells;
* a change of expected_route or should_escalate needs a note;
* every fixed case is still a valid TestCase.

Two sources of decisions:

* the review workbook (``read_rows``);
* Claude's first-pass CSV (``read_decisions_csv``) with the human audit applied on top
  (``apply_audit``): every audited row follows the auditor's verdict (an override on
  "disagree"); all cases get the tag ``ai_reviewed``, audited ones also
  ``human_verified``.

Review depth: priorities passed as ``bulk_priorities`` were accepted in bulk (marked keep
without a detailed look); their kept cases get the tag ``review_bulk_accepted`` so they can
be revisited, and the log records which priorities were fully reviewed and which weren't.

Cell meanings: a blank fixed_* cell = unchanged; fixed_expected_route "(none)" = route not
checked (null); fixed_should_escalate "(unset)" = not checked (null). Fixed cases keep their
id and get the tag ``review_fixed``. Output order is the dataset's order, so the result is
deterministic.
"""

import csv
import re
from collections import Counter
from collections.abc import Mapping, Sequence
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Final

from openpyxl import load_workbook
from pydantic import ValidationError

from evalkit.review.audit import AuditError, Verdict
from evalkit.review.opinions import tag_value
from evalkit.review.workbook import (
    DECISIONS,
    HEADERS,
    NO_ROUTE,
    REVIEW_SHEET,
    ROUTE_CHOICES,
    UNSET,
)
from evalkit.schemas import TestCase

FIXED_TAG: Final = "review_fixed"
BULK_TAG: Final = "review_bulk_accepted"
AI_TAG: Final = "ai_reviewed"
HUMAN_TAG: Final = "human_verified"
CSV_REQUIRED: Final = ("id", "decision", "note", "reviewer")
FIX_COLUMNS: Final = (
    "fixed_expected_route",
    "fixed_should_escalate",
    "fixed_question",
    "fixed_reference_answer",
    "remove_tags",  # optional (CSV only): tags to drop, separated by | or ,
    "fixed_gold_context_ids",  # optional (CSV only): replaces gold_context_ids (| or ,)
)


class ReviewError(Exception):
    """The workbook can't be applied; ``problems`` lists every reason."""

    def __init__(self, problems: Sequence[str]) -> None:
        super().__init__(f"{len(problems)} problem(s):\n" + "\n".join(f"- {p}" for p in problems))
        self.problems = list(problems)


@dataclass(frozen=True)
class Change:
    case_id: str
    generator: str
    field: str
    old: str
    new: str
    note: str


@dataclass
class ReviewResult:
    cases: list[TestCase] = field(default_factory=list)
    decisions: dict[str, str] = field(default_factory=dict)  # id -> keep/fix/delete
    generators: dict[str, str] = field(default_factory=dict)  # id -> generator
    changes: list[Change] = field(default_factory=list)
    deleted: list[tuple[str, str, str]] = field(default_factory=list)  # id, generator, note
    priorities: dict[str, int] = field(default_factory=dict)  # id -> review priority
    bulk_priorities: frozenset[int] = frozenset()


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):  # Excel may store a typed "true" as a boolean
        return "true" if value else "false"
    return str(value).strip()


def read_rows(path: Path) -> list[dict[str, str]]:
    ws = load_workbook(path, read_only=True, data_only=True)[REVIEW_SHEET]
    rows = ws.iter_rows(values_only=True)
    header = [_text(h) for h in next(rows)]
    missing = [h for h in HEADERS if h not in header]
    if missing:
        raise ReviewError([f"missing column(s): {', '.join(missing)}"])
    out = []
    for raw in rows:
        values = {h: _text(v) for h, v in zip(header, raw, strict=False)}
        if any(values.values()):
            out.append(values)
    return out


def _split(text: str) -> list[str]:
    return [t.strip() for t in re.split(r"[|,]", text) if t.strip()]


def _escalate_value(text: str) -> bool | None:
    return None if text == UNSET else text == "true"


def _show(value: object) -> str:
    if value is None:
        return "null"
    return str(value).lower() if isinstance(value, bool) else str(value)


def _priority(text: str) -> int | None:
    try:
        return int(float(text))
    except ValueError:
        return None


def apply_review(
    cases: Sequence[TestCase],
    rows: Sequence[dict[str, str]],
    bulk_priorities: frozenset[int] = frozenset(),
    extra_tags: Mapping[str, Sequence[str]] | None = None,
    chunk_ids: frozenset[str] | None = None,
) -> ReviewResult:
    """Apply decision rows (workbook or CSV form) to ``cases``. ``extra_tags`` adds tags
    per case id to every kept case."""
    by_id = {c.id: c for c in cases}
    problems: list[str] = []
    seen: Counter[str] = Counter(r.get("id", "") for r in rows)
    problems += [f"{i}: appears in {n} rows" for i, n in seen.items() if n > 1]
    problems += [f"{i}: not in the dataset" for i in seen if i not in by_id]
    problems += [f"{i}: no row in the workbook" for i in by_id if i not in seen]

    result = ReviewResult(bulk_priorities=bulk_priorities)
    edits: dict[str, dict[str, Any]] = {}
    removed: dict[str, list[str]] = {}
    for row in rows:
        cid = row.get("id", "")
        case = by_id.get(cid)
        if case is None:
            continue
        decision = row.get("decision", "").lower()
        note = row.get("notes", "")
        generator = tag_value(case, "gen-")
        result.generators[cid] = generator
        if bulk_priorities or row.get("priority", ""):
            prio = _priority(row.get("priority", ""))
            if prio is None:
                problems.append(f"{cid}: priority {row.get('priority', '')!r} is not a number")
                continue
            result.priorities[cid] = prio
        if decision not in DECISIONS:
            problems.append(
                f"{cid}: decision must be one of {', '.join(DECISIONS)}, got {decision!r}"
            )
            continue
        result.decisions[cid] = decision
        filled = [c for c in FIX_COLUMNS if row.get(c, "")]
        if decision != "fix":
            if filled:
                problems.append(f"{cid}: {', '.join(filled)} filled but decision is {decision!r}")
            if decision == "delete":
                result.deleted.append((cid, generator, note))
            continue
        if not filled:
            problems.append(f"{cid}: decision is 'fix' but no fixed_* cell is filled")
            continue
        update: dict[str, Any] = {}
        route = row.get("fixed_expected_route", "")
        if route:
            if route not in ROUTE_CHOICES:
                problems.append(f"{cid}: fixed_expected_route {route!r} is not a route")
                continue
            update["expected_route"] = None if route == NO_ROUTE else route
        esc = row.get("fixed_should_escalate", "").lower()
        if esc:
            if esc not in ("true", "false", UNSET):
                problems.append(f"{cid}: fixed_should_escalate {esc!r} is not true/false/(unset)")
                continue
            update["should_escalate"] = _escalate_value(esc)
        if row.get("fixed_question", ""):
            update["question"] = row["fixed_question"]
        if row.get("fixed_reference_answer", ""):
            update["reference_answer"] = row["fixed_reference_answer"]
        gold_ids = _split(row.get("fixed_gold_context_ids", ""))
        if gold_ids:
            unknown_ids = [i for i in gold_ids if chunk_ids is not None and i not in chunk_ids]
            if unknown_ids:
                problems.append(f"{cid}: fixed_gold_context_ids {unknown_ids} are not chunk ids")
                continue
            update["gold_context_ids"] = gold_ids
        drop = _split(row.get("remove_tags", ""))
        missing_tags = [t for t in drop if t not in case.tags]
        if missing_tags:
            problems.append(f"{cid}: remove_tags {missing_tags} not on the case")
            continue
        label_change = [
            k
            for k in ("expected_route", "should_escalate")
            if k in update and _show(update[k]) != _show(getattr(case, k))
        ]
        if label_change and not note:
            problems.append(f"{cid}: {' and '.join(label_change)} changed without a note")
            continue
        for key, new in update.items():
            old = getattr(case, key)
            if _show(new) != _show(old):
                result.changes.append(Change(cid, generator, key, _show(old), _show(new), note))
        if drop:
            result.changes.append(
                Change(cid, generator, "tags", ", ".join(drop), f"removed: {', '.join(drop)}", note)
            )
            removed[cid] = drop
        edits[cid] = update

    out: dict[str, TestCase] = {}
    for cid, decision in result.decisions.items():
        if decision == "delete":
            continue
        case = by_id[cid]
        update = edits.get(cid, {})
        new_tags = [FIXED_TAG] if cid in edits else []
        if result.priorities.get(cid) in bulk_priorities:
            new_tags.append(BULK_TAG)
        new_tags += (extra_tags or {}).get(cid, [])
        if not update and not new_tags:
            out[cid] = case
            continue
        data = case.model_dump(mode="json") | update
        kept_tags = [t for t in case.tags if t not in removed.get(cid, [])]
        data["tags"] = [*kept_tags, *(t for t in new_tags if t not in kept_tags)]
        try:
            out[cid] = TestCase.model_validate(data)
        except ValidationError as e:
            first = e.errors()[0]
            problems.append(f"{cid}: fixed case is invalid: {first['msg']}")
    if problems:
        raise ReviewError(problems)
    result.cases = [out[c.id] for c in cases if c.id in out]
    return result


def render_status(result: ReviewResult) -> list[str]:
    """Which priorities a human reviewed in full and which were accepted in bulk."""
    lines = [
        "## Review status",
        "",
        "| Priority | Rows | keep | fix | delete | How reviewed |",
        "|---|---|---|---|---|---|",
    ]
    for prio in sorted(set(result.priorities.values())):
        ids = [i for i, p in result.priorities.items() if p == prio]
        c = Counter(result.decisions[i] for i in ids)
        how = (
            f"bulk-accepted as keep without detailed review; kept cases tagged `{BULK_TAG}`"
            if prio in result.bulk_priorities
            else "fully reviewed by a human, row by row"
        )
        lines.append(f"| {prio} | {len(ids)} | {c['keep']} | {c['fix']} | {c['delete']} | {how} |")
    tagged = sum(BULK_TAG in c.tags for c in result.cases)
    lines += ["", f"Cases tagged `{BULK_TAG}`: {tagged}.", ""]
    return lines


def render_log(result: ReviewResult, method: Sequence[str] = ()) -> str:
    """The review log; ``method`` lines describe how the decisions were made."""
    per_gen: dict[str, Counter[str]] = {}
    for cid, decision in result.decisions.items():
        per_gen.setdefault(result.generators[cid], Counter())[decision] += 1
    lines = [
        f"Applied review: {len(result.cases)} cases kept "
        f"({sum(d == 'fix' for d in result.decisions.values())} of them fixed), "
        f"{len(result.deleted)} deleted.",
        "",
        *method,
        *(render_status(result) if result.bulk_priorities else []),
        "## Decisions per generator",
        "",
        "| Generator | keep | fix | delete |",
        "|---|---|---|---|",
    ]
    for gen in sorted(per_gen):
        c = per_gen[gen]
        lines.append(f"| {gen} | {c['keep']} | {c['fix']} | {c['delete']} |")
    total = Counter(result.decisions.values())
    lines.append(f"| **total** | {total['keep']} | {total['fix']} | {total['delete']} |")
    labels = [c for c in result.changes if c.field in ("expected_route", "should_escalate")]
    texts = [c for c in result.changes if c not in labels]
    lines += ["", f"## Label changes ({len(labels)})", ""]
    if labels:
        lines += ["| Case | Generator | Field | Old | New | Reason |", "|---|---|---|---|---|---|"]
        lines += [
            f"| `{c.case_id}` | {c.generator} | {c.field} | {c.old} | {c.new} | {_cell(c.note)} |"
            for c in labels
        ]
    lines += ["", f"## Other fixes: text, tags and gold context ({len(texts)})", ""]
    if texts:
        lines += ["| Case | Field | New text | Note |", "|---|---|---|---|"]
        lines += [
            f"| `{c.case_id}` | {c.field} | {_cell(c.new)} | {_cell(c.note)} |" for c in texts
        ]
    lines += ["", f"## Deleted ({len(result.deleted)})", ""]
    if result.deleted:
        lines += ["| Case | Generator | Note |", "|---|---|---|"]
        lines += [f"| `{i}` | {g} | {_cell(n) or '-'} |" for i, g, n in result.deleted]
    return "\n".join(lines) + "\n"


def _cell(text: str) -> str:
    return text.replace("|", "\\|").replace("\n", " ")


# --------------------------------------------------------------------------- CSV + audit


def read_decisions_csv(path: Path) -> list[dict[str, str]]:
    """Claude's decisions as decision rows (the CSV's ``note`` becomes ``notes``)."""
    with path.open(encoding="utf-8", newline="") as f:
        reader = csv.DictReader(f)
        missing = [c for c in CSV_REQUIRED if c not in (reader.fieldnames or [])]
        if missing:
            raise ReviewError([f"{path.name}: missing column(s): {', '.join(missing)}"])
        rows = []
        for raw in reader:
            row = {k: (v or "").strip() for k, v in raw.items() if k is not None}
            row["notes"] = row.pop("note", "")
            rows.append(row)
    problems = [f"{r['id']}: no reviewer" for r in rows if not r.get("reviewer")]
    if problems:
        raise ReviewError(problems)
    return rows


@dataclass(frozen=True)
class Override:
    case_id: str
    old_decision: str
    new_decision: str
    note: str


@dataclass
class AuditOutcome:
    rows: list[dict[str, str]]
    audited: list[str]
    overrides: list[Override]

    @property
    def agreed(self) -> int:
        return len(self.audited) - len(self.overrides)


def apply_audit(rows: Sequence[dict[str, str]], verdicts: Mapping[str, Verdict]) -> AuditOutcome:
    """Rows with the auditor's overrides applied (blank override fields = Claude's value)."""
    by_id = {r["id"]: r for r in rows}
    unknown = sorted(set(verdicts) - set(by_id))
    if unknown:
        raise AuditError([f"{i}: audited but not in the decisions" for i in unknown])
    out: list[dict[str, str]] = []
    overrides: list[Override] = []
    for row in rows:
        verdict = verdicts.get(row["id"])
        if verdict is None or verdict.agree:
            out.append(dict(row))
            continue
        o = verdict.override
        new = dict(row)
        new["decision"] = o["decision"]
        if o["decision"] == "fix":
            for key in ("fixed_expected_route", "fixed_should_escalate", "fixed_reference_answer"):
                if o.get(key):
                    new[key] = o[key]
        else:
            for key in ("fixed_expected_route", "fixed_should_escalate", "fixed_reference_answer"):
                new[key] = ""
        # keep Claude's note: other changes on the row (tags, gold context) still cite it
        earlier = f" Earlier: {row['notes']}" if row.get("notes") else ""
        new["notes"] = f"human override: {o['note']}.{earlier}"
        new["reviewer"] = "human"
        overrides.append(Override(row["id"], row["decision"], o["decision"], o["note"]))
        out.append(new)
    audited = [r["id"] for r in rows if r["id"] in verdicts]
    return AuditOutcome(out, audited, overrides)


def review_tags(outcome: AuditOutcome) -> dict[str, list[str]]:
    audited = set(outcome.audited)
    return {r["id"]: [AI_TAG, HUMAN_TAG] if r["id"] in audited else [AI_TAG] for r in outcome.rows}


def render_method(
    ai_rows: Sequence[dict[str, str]], outcome: AuditOutcome, seed_note: str
) -> list[str]:
    """The "Review method" section for the log: who decided what, and how it was checked."""
    ai = Counter(r["decision"] for r in ai_rows)
    changed = sum(r["decision"] != "keep" for r in ai_rows)
    sampled = len(outcome.audited) - sum(
        1 for r in ai_rows if r["id"] in set(outcome.audited) and r["decision"] != "keep"
    )
    lines = [
        "## Review method",
        "",
        f"1. **First pass by Claude (AI)** on all {len(ai_rows)} cases, against "
        "`agent_snapshot/SYSTEM_RULES.md` only (never the router's prediction): "
        f"keep {ai['keep']}, fix {ai['fix']}, delete {ai['delete']} "
        "(`review/ai_review_decisions.csv`, reviewer=claude).",
        f"2. **Human audit** (`review/human_audit.md`) of {len(outcome.audited)} cases: every "
        f"fix and delete ({changed}) plus {sampled} randomly sampled keeps ({seed_note}). "
        f"Agreed: {outcome.agreed}; overridden: {len(outcome.overrides)}.",
        f"3. The other {len(ai_rows) - len(outcome.audited)} keeps were **not checked by a "
        "human**. All cases are tagged `ai_reviewed`; the audited ones also `human_verified`.",
        "",
    ]
    if outcome.overrides:
        lines += ["| Overridden case | Claude | Human | Human's note |", "|---|---|---|---|"]
        lines += [
            f"| `{o.case_id}` | {o.old_decision} | {o.new_decision} | {_cell(o.note)} |"
            for o in outcome.overrides
        ]
        lines.append("")
    return lines
