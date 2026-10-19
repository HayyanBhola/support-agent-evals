"""The human review workbook for synthetic cases (``datasets/review_queue_v1.xlsx``).

Sheet "Review": one row per case, sorted by priority then generator (file order inside a
generator). Priorities:

1. router disagreement where the expected or router route is a safety route (escalate,
   medical, guardrail), or the case is safety_critical;
2. any other router disagreement;
3. remaining safety_critical cases;
4. everything else (agrees, or route not checked).

The reviewer fills ``decision`` (keep / fix / delete) and, for fixes, the ``fixed_*``
columns; ``scripts/apply_review.py`` applies them. Sheet "How to review" holds the rules.
"""

from collections.abc import Callable, Sequence
from dataclasses import dataclass
from pathlib import Path
from typing import Final

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from evalkit.review.opinions import AGREES, SUGGESTIONS, UNCHECKED, assess, tag_value
from evalkit.schemas import Route, TestCase

REVIEW_SHEET: Final = "Review"
HOWTO_SHEET: Final = "How to review"
SAFETY_ROUTES: Final = frozenset({"escalate", "medical", "guardrail"})
DECISIONS: Final = ("keep", "fix", "delete")
NO_ROUTE: Final = "(none)"
UNSET: Final = "(unset)"
ROUTE_CHOICES: Final = (*(r.value for r in Route), NO_ROUTE)
ESCALATE_CHOICES: Final = ("true", "false", UNSET)

# (column, width); order is the sheet's column order.
COLUMNS: Final[tuple[tuple[str, int], ...]] = (
    ("priority", 8),
    ("id", 24),
    ("generator", 12),
    ("evolution", 11),
    ("expected_route", 12),
    ("router_route", 12),
    ("your_opinion", 16),
    ("rule", 12),
    ("should_escalate", 10),
    ("safety_critical", 10),
    ("difficulty", 9),
    ("question", 50),
    ("reference_answer", 45),
    ("gold_facts", 32),
    ("gold_context_ids", 22),
    ("must_include", 22),
    ("must_include_any", 28),
    ("must_not_include", 28),
    ("decision", 10),
    ("fixed_expected_route", 14),
    ("fixed_should_escalate", 12),
    ("fixed_question", 40),
    ("fixed_reference_answer", 40),
    ("notes", 45),
)
HEADERS: Final = tuple(name for name, _ in COLUMNS)
DECISION_COLUMNS: Final = (
    "decision",
    "fixed_expected_route",
    "fixed_should_escalate",
    "fixed_question",
    "fixed_reference_answer",
    "notes",
)

HOWTO: Final = (
    ("How to review the synthetic cases", True),
    ("", False),
    ("Labels come from the agent's rules (agent_snapshot/SYSTEM_RULES.md, products.json, "
     "seed_data.sql, the knowledge-base chunks), NEVER from what the router or the agent does. "
     "router_route and your_opinion are review aids only.", False),
    ("", False),
    ("decision (required on every row):", True),
    ("keep: the case is right as it is (an agent-bug disagreement is usually a keep: the "
     "case is doing its job).", False),
    ("fix: the case is worth having but a label or the wording is wrong. Fill only the "
     "fixed_* cells you change; a blank fixed_* cell means unchanged.", False),
    ("delete: the case can't be made right with a small fix (unrealistic, ambiguous between "
     "two correct answers, duplicate in meaning, tests nothing useful).", False),
    ("", False),
    ("fixed_expected_route: a route, or (none) for 'route not checked'.", False),
    ("fixed_should_escalate: true / false, or (unset) for 'not checked'.", False),
    ("fixed_question / fixed_reference_answer: the corrected text. The case id stays the same.",
     False),
    ("", False),
    ("Every change of expected_route or should_escalate needs a note citing the SYSTEM_RULES "
     "rule (e.g. 'rule 9: reaction after use is a safety concern'). apply_review.py refuses "
     "to run without it, or if any row has no decision.", True),
    ("", False),
    ("Priorities: 1 = disagreement on a safety route (escalate, medical, guardrail) or a "
     "safety-critical case; 2 = other disagreements; 3 = remaining safety-critical cases; "
     "4 = everything else.", False),
    ("your_opinion: agrees / likely agent bug / likely label mistake / unclear / n/a (route "
     "not checked); reasons per case type are in notes/day5_generation.md.", False),
)  # fmt: skip


@dataclass(frozen=True)
class ReviewRow:
    priority: int
    generator: str
    order: int  # position in the dataset
    values: dict[str, object]


def _join(items: Sequence[str]) -> str:
    return " | ".join(items)


def compact_facts(facts: dict[str, object]) -> str:
    parts = []
    for key, value in facts.items():
        shown = ", ".join(value) if isinstance(value, list) else value
        parts.append(f"{key}={shown}")
    return "; ".join(parts)


def _bool(value: bool | None) -> str:
    return "" if value is None else str(value).lower()


def priority(case: TestCase, router_route: str, opinion: str) -> int:
    disagrees = opinion not in (AGREES, UNCHECKED)
    expected = case.expected_route.value if case.expected_route else None
    if disagrees and (
        expected in SAFETY_ROUTES or router_route in SAFETY_ROUTES or case.safety_critical
    ):
        return 1
    if disagrees:
        return 2
    return 3 if case.safety_critical else 4


def build_rows(cases: Sequence[TestCase], predict: Callable[[str], str]) -> list[ReviewRow]:
    rows = []
    for order, case in enumerate(cases):
        routed = predict(case.question)
        a = assess(case, routed)
        generator = tag_value(case, "gen-")
        values: dict[str, object] = {
            "priority": priority(case, routed, a.opinion),
            "id": case.id,
            "generator": generator,
            "evolution": tag_value(case, "evo-"),
            "expected_route": case.expected_route.value if case.expected_route else NO_ROUTE,
            "router_route": routed,
            "your_opinion": a.opinion,
            "rule": a.rule,
            "should_escalate": _bool(case.should_escalate),
            "safety_critical": _bool(case.safety_critical),
            "difficulty": case.difficulty.value,
            "question": case.question,
            "reference_answer": case.reference_answer or "",
            "gold_facts": compact_facts(dict(case.gold_db_facts)),
            "gold_context_ids": _join(case.gold_context_ids),
            "must_include": _join(case.must_include),
            "must_include_any": " ; ".join(_join(g) for g in case.must_include_any),
            "must_not_include": _join(case.must_not_include),
            **dict.fromkeys(DECISION_COLUMNS, ""),
            "notes": SUGGESTIONS.get(case.id, ""),
        }
        rows.append(ReviewRow(int(str(values["priority"])), generator, order, values))
    first_seen: dict[str, int] = {}
    for r in rows:
        first_seen.setdefault(r.generator, r.order)
    return sorted(rows, key=lambda r: (r.priority, first_seen[r.generator], r.order))


def write_workbook(rows: Sequence[ReviewRow], path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = REVIEW_SHEET
    ws.append(list(HEADERS))
    for row in rows:
        ws.append([row.values[h] for h in HEADERS])

    header_fill = PatternFill("solid", fgColor="DDEBF7")
    decision_fill = PatternFill("solid", fgColor="FFF2CC")
    wrap = Alignment(wrap_text=True, vertical="top")
    for idx, (name, width) in enumerate(COLUMNS, start=1):
        letter = get_column_letter(idx)
        ws.column_dimensions[letter].width = width
        head = ws.cell(row=1, column=idx)
        head.font = Font(bold=True)
        head.fill = decision_fill if name in DECISION_COLUMNS else header_fill
        head.alignment = wrap
        for cell in ws[letter][1:]:
            cell.alignment = wrap
    ws.freeze_panes = "C2"  # header row and the priority/id columns stay visible
    ws.auto_filter.ref = ws.dimensions

    last = len(rows) + 1
    for name, choices in (
        ("decision", DECISIONS),
        ("fixed_expected_route", ROUTE_CHOICES),
        ("fixed_should_escalate", ESCALATE_CHOICES),
    ):
        letter = get_column_letter(HEADERS.index(name) + 1)
        dv = DataValidation(type="list", formula1=f'"{",".join(choices)}"', allow_blank=True)
        dv.error = f"Choose one of: {', '.join(choices)}"
        dv.showErrorMessage = True
        ws.add_data_validation(dv)
        dv.add(f"{letter}2:{letter}{last}")

    howto = wb.create_sheet(HOWTO_SHEET)
    howto.column_dimensions["A"].width = 110
    for text, bold in HOWTO:
        howto.append([text])
        cell = howto.cell(row=howto.max_row, column=1)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.font = Font(bold=bold)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
